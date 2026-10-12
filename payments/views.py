from io import BytesIO
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.http import HttpResponse, Http404
from django.db.models import Q

from accounts.decorators import role_required
from accounts.permissions import reviewable_stages, can_edit_record, is_personal_view
from banners import workflow as proposal_workflow
from banners.models import Store
from .models import PaymentClaim
from .forms import PaymentClaimForm
from . import workflow


APPROVAL_FILL_COLORS = {
    'in_review': 'FFF3CD',
    'rework': 'CFE2FF',
    'approved': 'D1E7DD',
    'rejected': 'F8D7DA',
}


def _parse_date(value):
    if not value:
        return None
    try:
        return datetime.strptime(str(value).strip(), '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return None


def _detail_queryset():
    return PaymentClaim.objects.select_related(
        'created_by', 'work_order', 'work_order__proposal', 'work_order__proposal__campaign',
        'work_order__proposal__vendor', 'work_order__proposal__category', 'work_order__proposal__department',
        'current_step__role',
    )


def _filtered_payments(request):
    claims = _detail_queryset()
    if is_personal_view(request.user):
        claims = claims.filter(created_by=request.user)

    search = request.GET.get('q', '').strip()
    if search:
        claims = claims.filter(
            Q(payment_number__icontains=search) |
            Q(work_order__po_number__icontains=search) |
            Q(work_order__proposal__proposal_number__icontains=search) |
            Q(work_order__proposal__vendor__name__icontains=search) |
            Q(work_order__proposal__campaign__name__icontains=search) |
            Q(work_order__proposal__location__icontains=search) |
            Q(created_by__username__icontains=search)
        )

    status = request.GET.get('status', '').strip()
    if status and status != 'all':
        claims = claims.filter(approval_status=status)

    date_range = request.GET.get('date_range', '').strip()
    if date_range:
        parts = date_range.split(' to ')
        if len(parts) == 2:
            d_start = _parse_date(parts[0])
            d_end = _parse_date(parts[1])
            if d_start:
                claims = claims.filter(created_at__date__gte=d_start)
            if d_end:
                claims = claims.filter(created_at__date__lte=d_end)
        elif len(parts) == 1:
            d_start = _parse_date(parts[0])
            if d_start:
                claims = claims.filter(created_at__date=d_start)

    date_from = request.GET.get('date_from', '').strip()
    if date_from:
        d_from = _parse_date(date_from)
        if d_from:
            claims = claims.filter(created_at__date__gte=d_from)

    store = request.GET.get('store', '').strip()
    if store:
        claims = claims.filter(
            Q(work_order__proposal__location__iexact=store) |
            Q(work_order__proposal__items__store_name__iexact=store)
        ).distinct()

    date_to = request.GET.get('date_to', '').strip()
    if date_to:
        d_to = _parse_date(date_to)
        if d_to:
            claims = claims.filter(created_at__date__lte=d_to)


    sort = request.GET.get('sort', '-created_at').strip()
    if sort not in ('-created_at', 'created_at', 'payment_number', '-payment_number'):
        sort = '-created_at'
    claims = claims.order_by(sort)

    return claims


@role_required('can_add_payment')
def payment_upload(request):
    """Payments Upload page — submit bill/work order/proof/covering letter
    against an approved Work Order. Same submit-for-approval flow as
    adding a Proposal."""
    if request.method == 'POST':
        form = PaymentClaimForm(request.POST, request.FILES)
        if form.is_valid():
            claim = form.save(commit=False)
            claim.created_by = request.user
            claim.save()
            workflow.submit(claim, request.user)
            proposal_workflow.invalidate_pending_counts_for_step(claim.current_step)
            messages.success(request, f'Payment {claim.payment_number} submitted for approval.')
            return redirect('payment_detail', pk=claim.pk)
    else:
        form = PaymentClaimForm()

    return render(request, 'payments/payment_upload.html', {
        'form': form,
        'next_payment_number': PaymentClaim.next_payment_number_preview(),
    })


@role_required('can_view_payments')
def payment_list(request):
    """Payments Records page — every payment claim submitted so far.

    Same ownership rule as the Dashboard/Proposals list: a plain
    Executive only sees payments they created themselves; anyone with
    broader responsibility (reviewer/admin/superuser) sees every payment,
    org-wide."""
    claims = _filtered_payments(request)
    querydict = request.GET.copy()
    querydict.pop('page', None)
    querystring = querydict.urlencode()
    page_obj = Paginator(claims, 25).get_page(request.GET.get('page'))
    return render(request, 'payments/payment_list.html', {
        'claims': page_obj,
        'page_obj': page_obj,
        'approval_choices': PaymentClaim.APPROVAL_CHOICES,
        'store_choices': Store.objects.all(),
        'querystring': querystring,
    })



@role_required('can_view_payments')
def export_payments_excel(request):
    """Exports payment claims matching the active filters as an Excel (.xlsx) file."""
    claims = _filtered_payments(request)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Payments Records'

    headers = [
        'Payment #', 'Work Order #', 'Proposal #', 'Vendor',
        'Campaign', 'Department', 'Stores / Location', 'Submitted By',
        'Submission Date', 'Status', 'Remarks'
    ]
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    row_index = 2
    for claim in claims:
        prop = getattr(claim.work_order, 'proposal', None) if claim.work_order else None
        sheet.cell(row=row_index, column=1, value=claim.payment_number or '—')
        sheet.cell(row=row_index, column=2, value=claim.work_order.po_number if claim.work_order else '—')
        sheet.cell(row=row_index, column=3, value=prop.proposal_number if prop else '—')
        sheet.cell(row=row_index, column=4, value=prop.vendor.name if prop and prop.vendor else '—')
        sheet.cell(row=row_index, column=5, value=prop.campaign.name if prop and prop.campaign else '—')
        sheet.cell(row=row_index, column=6, value=prop.department.name if prop and prop.department else '—')
        sheet.cell(row=row_index, column=7, value=prop.location if prop else '—')
        sheet.cell(row=row_index, column=8, value=claim.created_by.username if claim.created_by else '—')
        sheet.cell(row=row_index, column=9, value=claim.created_at.strftime('%Y-%m-%d %H:%M') if claim.created_at else '—')

        status_cell = sheet.cell(row=row_index, column=10, value=claim.stage_label)
        status_cell.fill = PatternFill(
            start_color=APPROVAL_FILL_COLORS.get(claim.approval_status, 'FFFFFF'),
            end_color=APPROVAL_FILL_COLORS.get(claim.approval_status, 'FFFFFF'),
            fill_type='solid',
        )
        sheet.cell(row=row_index, column=11, value=claim.remarks or '')
        row_index += 1

    column_widths = [18, 22, 18, 20, 20, 18, 22, 16, 20, 22, 30]
    for col_index, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    sheet.freeze_panes = 'A2'

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="payments_records.xlsx"'
    return response


@role_required('can_edit_payment')
def payment_edit(request, pk):
    claim = get_object_or_404(PaymentClaim, pk=pk)
    if not can_edit_record(request.user, claim):
        raise PermissionDenied(
            "You can only edit payments you created, and only before they're approved."
        )
    needs_rework = claim.approval_status == 'rework'

    if request.method == 'POST':
        form = PaymentClaimForm(request.POST, request.FILES, instance=claim)
        if form.is_valid():
            form.save()
            if needs_rework:
                # Editing a reworked claim resubmits it — back to the
                # first approval step.
                workflow.resubmit(claim, request.user)
                proposal_workflow.invalidate_pending_counts_for_step(claim.current_step)
                messages.success(request, f'Payment {claim.payment_number} resubmitted for approval.')
            else:
                messages.success(request, 'Payment updated successfully.')
            return redirect('payment_detail', pk=claim.pk)
        messages.error(request, 'Please correct the errors below.')
    else:
        form = PaymentClaimForm(instance=claim)

    # (field name, bound field, currently-stored file or None) for each
    # upload — lets the edit template show a download link for what's
    # already uploaded, next to the "replace it" file input.
    documents = [
        (name, form[name], getattr(claim, name) if getattr(claim, name) else None)
        for name in ('bill_file', 'work_order_file', 'proof_file', 'covering_letter_file')
    ]

    return render(request, 'payments/payment_edit.html', {
        'form': form,
        'claim': claim,
        'documents': documents,
        'needs_rework': needs_rework,
        'rework_note': claim.approval_history.filter(action='rework').order_by('-created_at').first() if needs_rework else None,
    })


@role_required('can_delete_payment')
def payment_delete(request, pk):
    claim = get_object_or_404(_detail_queryset(), pk=pk)
    if request.method == 'POST':
        claim.delete()
        messages.success(request, 'Payment claim deleted successfully.')
        return redirect('payment_list')
    return render(request, 'payments/payment_confirm_delete.html', {'claim': claim})


@role_required('can_view_payments')
def payment_detail(request, pk):
    claim = get_object_or_404(_detail_queryset(), pk=pk)
    return render(request, 'payments/payment_detail.html', {
        'claim': claim,
        'history': claim.approval_history.select_related('actor').all(),
        'my_review_stages': reviewable_stages(request.user),
    })


@login_required
def payment_approvals(request):
    """Shared review queue: every payment claim currently sitting at any
    step this user is permitted to act on."""
    from accounts.permissions import can_review_anything
    if not can_review_anything(request.user):
        raise PermissionDenied("You don't have permission to review payments.")

    steps = reviewable_stages(request.user)
    if request.user.is_superuser:
        claims = PaymentClaim.objects.filter(approval_status='in_review')
    elif steps:
        claims = PaymentClaim.objects.filter(
            approval_status='in_review',
            current_step_id__in=steps,
        )
    else:
        claims = PaymentClaim.objects.none()

    claims = claims.select_related(
        'created_by', 'work_order', 'work_order__proposal', 'work_order__proposal__campaign', 'current_step__role'
    ).order_by('current_step__order', '-created_at')
    page_obj = Paginator(claims, 25).get_page(request.GET.get('page'))
    return render(request, 'payments/payment_approvals.html', {'claims': page_obj, 'page_obj': page_obj})


@login_required
def payment_review(request, pk):
    """The summary-before-approval screen for a payment claim - full
    details, complete approval history, and step-appropriate action
    buttons with a comment box. Same layout as proposal_review."""
    claim = get_object_or_404(_detail_queryset(), pk=pk)
    step = claim.current_step
    my_steps = reviewable_stages(request.user)

    can_act = (step is not None) and ((step.id in my_steps) or request.user.is_superuser)
    can_rework = bool(step) and step.allow_rework
    can_reject = bool(step) and getattr(step, 'can_reject', step.order > 2)
    has_next_step = bool(step) and proposal_workflow.next_step(step) is not None
    is_first_step = bool(step) and proposal_workflow.previous_step(step) is None
    is_next_level = bool(step) and step.order in (1, 2)
    is_level_3 = bool(step) and step.order == 3
    is_final_step = bool(step) and (step.order > 3 or proposal_workflow.next_step(step) is None)

    return render(request, 'payments/payment_review.html', {
        'claim': claim,
        'history': claim.approval_history.select_related('actor').all(),
        'can_act': can_act,
        'can_rework': can_rework,
        'can_reject': can_reject,
        'stage_label': claim.stage_label,
        'is_first_step': is_first_step,
        'has_next_step': has_next_step,
        'is_next_level': is_next_level,
        'is_level_3': is_level_3,
        'is_final_step': is_final_step,
    })


@login_required
def payment_approve_action(request, pk, action):
    if action not in ('approve', 'reject', 'rework'):
        raise Http404
    claim = get_object_or_404(PaymentClaim.objects.select_related('current_step__role'), pk=pk)
    step = claim.current_step

    if step is None:
        raise PermissionDenied("This payment is not currently in review.")

    if step.id not in reviewable_stages(request.user) and not request.user.is_superuser:
        raise PermissionDenied("You don't have permission to act on this payment at its current stage.")

    if request.method == 'POST':
        comment = request.POST.get('comment', '').strip()
        if not comment:
            messages.error(request, 'A comment is required for every decision.')
            return redirect('payment_review', pk=claim.pk)
        if action == 'rework' and not step.allow_rework:
            raise PermissionDenied("This stage has no rework option.")
        if action == 'reject' and not getattr(step, 'can_reject', step.order > 2):
            raise PermissionDenied("Step 1 and Step 2 do not have a reject option.")

        workflow.apply_action(claim, step, action, request.user, comment=comment)

        if action == 'approve':
            if step.order in (1, 2):
                upcoming = claim.current_step
                verb = f'sent to the next level ({upcoming.role.name} review)' if upcoming else 'sent to the next level'
            elif step.order == 3:
                upcoming = claim.current_step
                verb = f'approved and sent to {upcoming.role.name} review' if upcoming else 'fully approved'
            else:
                verb = 'fully approved'
        else:
            verb = {'reject': 'rejected', 'rework': 'sent for rework'}[action]
        messages.success(request, f'Payment {claim.payment_number} {verb}.')

        proposal_workflow.invalidate_pending_counts_for_step(step)
        proposal_workflow.invalidate_pending_counts_for_step(claim.current_step)
    return redirect('payment_approvals')
